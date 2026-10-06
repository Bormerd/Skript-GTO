from pathlib import Path
from openpyxl import load_workbook, Workbook
from difflib import SequenceMatcher
import re

BASE = Path(".")
COLLEGES_FOLDER = BASE / "Колледжи"

def normalize_uin(value):
    if value is None: return ""
    s = str(value).strip().replace(" ", "").replace("\u00a0", "")
    s = s.replace("-", "").replace("–", "").replace("—", "")
    s = s.replace("\n", "").replace("\r", "")
    if s.endswith(".0"): s = s[:-2]
    return s

def normalize_fio(value):
    if value is None: return ""
    s = str(value).strip().lower().replace("ё", "е")
    s = re.sub(r"[.,;:]+", " ", s)
    return " ".join(s.split())

def header_key(value):
    if value is None: return ""
    return re.sub(r"[^а-яa-z0-9]", "", str(value).lower().replace("ё", "е"))

def find_columns(ws):
    uin_col = fio_col = header_row = None
    for row in ws.iter_rows(max_row=min(60, ws.max_row)):
        found = False
        for cell in row:
            key = header_key(cell.value)
            if not key: continue
            if "уин" in key or ("идентификатор" in key and "участ" in key):
                uin_col, header_row, found = cell.column, cell.row, True
            if key in ("фио", "фиоучастника", "фамилияимяотчество") or "фио" in key:
                fio_col, header_row, found = cell.column, cell.row, True
        if found and uin_col: break
    return uin_col, fio_col, header_row

def find_main_file():
    candidates = []
    for p in BASE.glob("*.xlsx"):
        try:
            wb = load_workbook(p, read_only=True, data_only=True)
            sheets = set(wb.sheetnames)
            wb.close()
            if {"Золото", "Серебро", "Бронза"}.issubset(sheets):
                candidates.append(p)
        except Exception:
            pass
    if not candidates:
        raise FileNotFoundError("Не найден основной XLSX с листами Золото/Серебро/Бронза.")
    return candidates[0]

def main():
    main_file = find_main_file()
    uin_map, fio_map = {}, {}
    errors = []

    for file in COLLEGES_FOLDER.rglob("*.xlsx"):
        college = file.parent.name
        try:
            wb = load_workbook(file, data_only=True, read_only=True)
            for ws in wb.worksheets:
                ucol, fcol, hrow = find_columns(ws)
                if not ucol or not hrow: continue
                for r in range(hrow + 1, ws.max_row + 1):
                    u = normalize_uin(ws.cell(r, ucol).value)
                    f = normalize_fio(ws.cell(r, fcol).value) if fcol else ""
                    if u: uin_map.setdefault(u, set()).add(college)
                    if f: fio_map.setdefault(f, []).append((u, college, file.name))
            wb.close()
        except Exception as e:
            errors.append((str(file), str(e)))

    wb = load_workbook(main_file)
    total = by_uin = by_fio = not_found = conflicts = 0
    diagnostics = []

    for sheet in ("Золото", "Серебро", "Бронза"):
        if sheet not in wb.sheetnames: continue
        ws = wb[sheet]
        ucol, fcol, hrow = find_columns(ws)
        if not ucol or not hrow: continue

        college_col = next((c for c in range(1, ws.max_column+1)
                            if normalize_fio(ws.cell(hrow,c).value) == "колледж"), None)
        if college_col is None:
            college_col = ws.max_column + 1
            ws.cell(hrow, college_col, "Колледж")

        for r in range(hrow + 1, ws.max_row + 1):
            u, f = normalize_uin(ws.cell(r, ucol).value), normalize_fio(ws.cell(r, fcol).value) if fcol else ""
            if not u and not f: continue
            total += 1

            uc = uin_map.get(u, set())
            fh = fio_map.get(f, [])
            fcols = sorted(set(x[1] for x in fh if x[1]))

            if len(uc) == 1:
                college, status = next(iter(uc)), "УИН"
                by_uin += 1
            elif len(uc) > 1:
                college, status = "КОНФЛИКТ: несколько колледжей", "КОНФЛИКТ"
                conflicts += 1
            elif len(fcols) == 1:
                college, status = fcols[0], "ФИО"
                by_fio += 1
            elif len(fcols) > 1:
                college, status = "ТРЕБУЕТ ПРОВЕРКИ", "ФИО-КОНФЛИКТ"
                conflicts += 1
            else:
                college, status = "НЕ НАЙДЕН", "НЕ НАЙДЕН"
                not_found += 1

            ws.cell(r, college_col, f"{college} ({status})")
            if status != "УИН":
                diagnostics.append((sheet, f, u, status))

    out = BASE / f"{main_file.stem}_с_колледжами.xlsx"
    wb.save(out)

    rep = Workbook()
    wr = rep.active
    wr.title = "Диагностика"
    wr.append(["Лист","ФИО","УИН","Статус"])
    for row in diagnostics: wr.append(row)
    ws = rep.create_sheet("Статистика")
    for row in [
        ("Основной файл", main_file.name),
        ("УИН в индексе", len(uin_map)),
        ("ФИО в индексе", len(fio_map)),
        ("Всего участников", total),
        ("По УИН", by_uin),
        ("По ФИО", by_fio),
        ("Не найдено", not_found),
        ("Конфликты", conflicts),
        ("Ошибки чтения", len(errors)),
    ]: ws.append(row)
    rep.save(BASE / f"{main_file.stem}_ДИАГНОСТИКА.xlsx")

    print(f"Основной файл: {main_file.name}")
    print(f"Всего: {total}; по УИН: {by_uin}; по ФИО: {by_fio}; не найдено: {not_found}; конфликтов: {conflicts}")

if __name__ == "__main__":
    main()
